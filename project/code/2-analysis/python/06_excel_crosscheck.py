"""Langkah 7 (bagian 2, Q14): cek silang KPI SQL dengan rumus Excel.

Workbook berisi data bersih dan sheet "cek" dengan rumus SUMIFS/COUNTIFS/PRODUCT. Nilai SQL ditulis
sebagai angka di sebelahnya. Excel (lewat COM, Windows) menghitung ulang rumus, lalu hasilnya dibaca kembali.
Hasil: results/2-analysis/crosscheck.xlsx dan crosscheck_excel.csv (semua baris harus PASS).
Tanpa Excel, workbook tetap dibuat tetapi status tidak bisa diisi (skrip berhenti dengan pesan).
"""
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

import duckdb
import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

sys.path.insert(0, str(Path(__file__).resolve().parents[2]))
from common import CLEAN_DIR, DB_PATH, RES_ANALYSIS, STATIONS, result

prod = pd.read_csv(CLEAN_DIR / "production_daily.csv")
dlog = pd.read_csv(CLEAN_DIR / "defect_log.csv")
rlog = pd.read_csv(CLEAN_DIR / "rework_log.csv")
st = pd.read_csv(result("station.csv")).set_index("station")
line = pd.read_csv(result("line_kpi.csv")).iloc[0]
par = pd.read_csv(result("pareto_line.csv"))
sh = pd.read_csv(result("shift_station.csv")).set_index("station")
rd = pd.read_csv(result("rework_defect.csv")).set_index("defect_code")
con = duckdb.connect(str(DB_PATH), read_only=True)

wb = Workbook()
ws = wb.active
ws.title = "cek"
prod["date"] = pd.to_datetime(prod.date).dt.date           # tanggal asli Excel, bukan teks
for name, df in (("produksi", prod[["date", "shift", "station", "model", "input_qty", "ok_first_pass", "ng_qty"]]),
                 ("defect", dlog[["date", "shift", "station", "defect_code", "qty"]]),
                 ("rework", rlog[["station_origin", "defect_code", "result"]].fillna(""))):
    s = wb.create_sheet(name)
    s.append(list(df.columns))
    for row in df.itertuples(index=False):
        s.append(list(row))
NP, ND, NR = len(prod) + 1, len(dlog) + 1, len(rlog) + 1
P = lambda c: f"produksi!${c}$2:${c}${NP}"
D = lambda c: f"defect!${c}$2:${c}${ND}"
R = lambda c: f"rework!${c}$2:${c}${NR}"

ws.append(["KPI", "Rumus Excel", "Nilai SQL", "Selisih", "Status"])
rows = []                                            # (label, rumus, nilai SQL)
for i, s in enumerate(STATIONS):
    rows.append((f"Input | {s}", f'=SUMIFS({P("E")},{P("C")},"{s}")', st.input_qty[s]))
for i, s in enumerate(STATIONS):
    rows.append((f"FPY | {s}", f'=SUMIFS({P("F")},{P("C")},"{s}")/SUMIFS({P("E")},{P("C")},"{s}")', st.fpy[s]))
fpy0 = 2 + len(STATIONS)                             # baris pertama FPY (baris 2-11 = input)
for s in STATIONS:
    rows.append((f"Final yield | {s}",
                 f'=(SUMIFS({P("F")},{P("C")},"{s}")+COUNTIFS({R("A")},"{s}",{R("C")},"OK"))/SUMIFS({P("E")},{P("C")},"{s}")',
                 st.final_yield[s]))
fy0 = fpy0 + len(STATIONS)
rows.append(("RTY (hasil kali FPY)", f"=PRODUCT(B{fpy0}:B{fpy0 + 9})", line.rty))
rows.append(("Final yield lini", f"=PRODUCT(B{fy0}:B{fy0 + 9})", line.final_yield))
rows.append(("Rework rate lini", f'=SUM({P("G")})/B2', line.rework_rate))
rows.append(("Scrap rate lini", f'=COUNTIF({R("C")},"Scrap")/B2', line.scrap_rate))
rows.append(("Recovery rate", f'=COUNTIF({R("C")},"OK")/(COUNTIF({R("C")},"OK")+COUNTIF({R("C")},"Scrap"))', line.recovery_rate))
for r in par.head(3).itertuples():
    rows.append((f"Pareto share | {r.defect_code}", f'=SUMIFS({D("E")},{D("D")},"{r.defect_code}")/SUM({D("E")})', r.share))
s = "Component Placement"
rows.append((f"Selisih NG rate Shift 2 - Shift 1 | {s}",
             f'=SUMIFS({P("G")},{P("C")},"{s}",{P("B")},"Shift 2")/SUMIFS({P("E")},{P("C")},"{s}",{P("B")},"Shift 2")'
             f'-SUMIFS({P("G")},{P("C")},"{s}",{P("B")},"Shift 1")/SUMIFS({P("E")},{P("C")},"{s}",{P("B")},"Shift 1")',
             sh.gap[s]))
rows.append(("Recovery rate | SCR-02",
             f'=COUNTIFS({R("B")},"SCR-02",{R("C")},"OK")/(COUNTIFS({R("B")},"SCR-02",{R("C")},"OK")+COUNTIFS({R("B")},"SCR-02",{R("C")},"Scrap"))',
             rd.recovery_rate["SCR-02"]))
for day in ("2026-03-02", "2026-04-08", "2026-05-20"):          # sampel hari x stasiun
    for s in ("Component Placement", "Reflow Soldering", "Screen Assembly"):
        v = con.execute("SELECT sum(ok_first_pass) / sum(input_qty) FROM production_daily WHERE date = ? AND station = ?",
                        [day, s]).fetchone()[0]
        xd = "DATE({},{},{})".format(*[int(t) for t in day.split("-")])
        rows.append((f"FPY sampel | {day} | {s}",
                     f'=SUMIFS({P("F")},{P("A")},{xd},{P("C")},"{s}")/SUMIFS({P("E")},{P("A")},{xd},{P("C")},"{s}")', v))
con.close()
for k, (label, f, v) in enumerate(rows, start=2):
    ws.append([label, f, float(v), f"=ABS(B{k}-C{k})", f'=IF(D{k}<1E-9,"PASS","FAIL")'])
ws.column_dimensions["A"].width = 52
for c in "BCD":
    ws.column_dimensions[c].width = 16
for cell in ws[1]:
    cell.font = Font(bold=True, color="FFFFFF")
    cell.fill = PatternFill("solid", fgColor="1F3A5F")

out = RES_ANALYSIS / "crosscheck.xlsx"
tmp = Path(tempfile.mkdtemp())                     # jalur pendek: Excel menolak jalur > 218 karakter
x = tmp / "crosscheck.xlsx"
wb.save(x)
ps = (f"$x=New-Object -ComObject Excel.Application;$x.DisplayAlerts=$false;$w=$x.Workbooks.Open('{x}');"
      "$x.CalculateFull();$w.Save();$w.Close();$x.Quit()")
try:
    subprocess.run(["powershell", "-NoProfile", "-Command", ps], check=True, timeout=300)
except (OSError, subprocess.SubprocessError) as e:
    shutil.copy(x, out)
    sys.exit(f"Excel tidak tersedia ({e}); workbook disimpan tanpa hasil hitung ulang.")
shutil.copy(x, out)
shutil.rmtree(tmp, ignore_errors=True)

vals = load_workbook(out, data_only=True, read_only=True)["cek"]
res = pd.DataFrame([r[:5] for r in vals.iter_rows(min_row=2, values_only=True)],
                   columns=["kpi", "excel", "sql", "abs_diff", "status"])
res.to_csv(RES_ANALYSIS / "crosscheck_excel.csv", index=False, float_format="%.10g")
print(res.status.value_counts().to_string())
assert (res.status == "PASS").all(), res[res.status != "PASS"]
